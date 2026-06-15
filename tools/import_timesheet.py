#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Конвертер годового табеля (лист «2026» Excel) в формат приложения Фотомеханика.

Источник — лист «2026», где каждая строка самодостаточна (с row 19):
    A=Месяц («Январь 2026»)  B=График  C=ФИО  D=Проект  E=Признак времени
    F..(6-36)=Календарные дни 1-31   37=ИТОГО
(Месячные вкладки янв/мар/… — старые данные 2023 года, не используются.)

Привязка к сотрудникам — по display_name из users.json (точное совпадение с ФИО),
с фолбэком по «фамилия + первая буква имени».

Правила переноса (по договорённости):
  • числовые часы → строки отчётов (лист «Отчёты»);
  • признак времени «Обед» → пропускаем (приложение само добавляет обед);
  • проект «Отпуск» → отпуск (зелёный) в vacations.json (по user_id), без часов;
  • признак «Переработки» → жёлтый флаг, «Отработка невыхода» → синий флаг
    (report_flags.json: {uid: {date: {project: overtime|dayoff}}});
  • комментарии не переносятся.

Выход (каталог --out):
  reports_import.tsv      — строки для листа «Отчёты» (сопоставленные сотрудники)
  reports_unmatched.tsv   — то же для не найденных в users.json (на всякий случай)
  vacations_import.json    — отпуска по user_id
  report_flags_import.json — флаги переработка/отработка по user_id
  unmatched.txt            — ФИО, не найденные в users.json

Запуск:
  python3 tools/import_timesheet.py --xlsx табель.xlsx --users users.json --out import_out
"""
import argparse
import json
import os
import re
import uuid
import collections

import openpyxl

SHEET = '2026'
DATA_START_ROW = 19
DAY_HEADER_ROW = 17
COL_MONTH, COL_FIO, COL_PROJECT, COL_PRIZNAK = 1, 3, 4, 5

MONTH_RU = {
    'январь': 1, 'февраль': 2, 'март': 3, 'апрель': 4, 'май': 5, 'июнь': 6,
    'июль': 7, 'август': 8, 'сентябрь': 9, 'октябрь': 10, 'ноябрь': 11, 'декабрь': 12,
}
SKIP_PRIZNAK = {'обед'}
FLAG_MAP = {'переработки': 'overtime', 'отработка невыхода': 'dayoff'}

SUR_SUF = re.compile(r'(ов|ёв|ев|ин|ын|ский|цкий|ской|их|ых|ко|юк|ук|ян|дзе|швили)$', re.I)
PATR    = re.compile(r'(ович|евич|ьич|инич|овна|евна|ична)$', re.I)


def name_key(raw: str):
    toks = raw.split()
    sur = next((t for t in toks if not PATR.search(t) and SUR_SUF.search(t)), None) or toks[0]
    given = next((t for t in toks if t != sur and not PATR.search(t)), toks[-1])
    return (sur.lower(), given[:1].lower())


def parse_month(a):
    """«Январь 2026» → (2026, 1)."""
    m = re.match(r'\s*([А-Яа-яЁё]+)\s+(\d{4})', str(a or ''))
    if not m:
        return None
    mon = MONTH_RU.get(m.group(1).lower())
    return (int(m.group(2)), mon) if mon else None


def day_columns(ws):
    cols = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(DAY_HEADER_ROW, c).value
        if isinstance(v, (int, float)) and 1 <= int(v) <= 31 and float(v) == int(v):
            cols[c] = int(v)
    return cols


def fmt_hours(v) -> str:
    return str(int(v)) if float(v) == int(v) else str(v).replace('.', ',')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', required=True)
    ap.add_argument('--users', help='users.json для привязки ФИО к username/id')
    ap.add_argument('--out', default='import_out')
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb[SHEET]
    dcols = day_columns(ws)

    users = json.load(open(args.users, encoding='utf-8')) if args.users and os.path.exists(args.users) else {}
    by_display, by_key = {}, {}
    for uid, u in users.items():
        dn = (u.get('display_name') or u.get('username') or '').strip()
        if not dn:
            continue
        rec = {'username': u.get('username', ''), 'id': str(uid)}
        by_display[dn.lower()] = rec
        by_key.setdefault(name_key(dn), rec)

    canon = {}
    matched_keys = set()

    def resolve(fio):
        k = name_key(fio)
        if len(fio) > len(canon.get(k, '')):
            canon[k] = fio
        rec = by_display.get(fio.lower()) or by_key.get(k)
        if rec:
            matched_keys.add(k)
            return rec['username'], rec['id'], True
        return canon[k], '', False

    reports = []                                   # (date, username, project, hours, matched)
    vacations = collections.defaultdict(dict)      # uid → {date: True}
    flags = collections.defaultdict(lambda: collections.defaultdict(dict))  # uid → date → {project: type}
    work_days = collections.defaultdict(set)
    vac_names = collections.defaultdict(set)
    months_seen = collections.Counter()
    skipped_lunch = 0

    for r in range(DATA_START_ROW, ws.max_row + 1):
        fio = ws.cell(r, COL_FIO).value
        if not (fio and isinstance(fio, str) and fio.strip()) or fio.strip() == 'ФИО' or fio.strip().endswith('ИТОГ'):
            continue
        fio = fio.strip()
        ym = parse_month(ws.cell(r, COL_MONTH).value)
        if not ym:
            continue
        year, month = ym
        months_seen[f"{year}-{month:02d}"] += 1
        project = str(ws.cell(r, COL_PROJECT).value or '').strip()
        priznak = str(ws.cell(r, COL_PRIZNAK).value or '').strip().lower()

        username, uid, matched = resolve(fio)
        is_vacation = 'отпуск' in project.lower()
        flag_type = FLAG_MAP.get(priznak)

        for c, day in dcols.items():
            v = ws.cell(r, c).value
            if v is None or v == '':
                continue
            date_iso = f"{year:04d}-{month:02d}-{day:02d}"
            if priznak in SKIP_PRIZNAK:            # обед — приложение добавит сам
                skipped_lunch += 1
                continue
            if is_vacation:
                if uid:
                    vacations[uid][date_iso] = True
                vac_names[username].add(date_iso)
                continue
            if not isinstance(v, (int, float)) or not v:
                continue                            # буквенные пометки в часах — пропуск
            reports.append((date_iso, username, project, float(v), matched))
            work_days[(username, date_iso)].add(project)
            if flag_type and uid:
                flags[uid][date_iso][project] = flag_type

    # Конфликт «отпуск + часы в один день» → приоритет у часов, отпуск снимаем
    conflicts = 0
    for uid, days in list(vacations.items()):
        uname = next((u['username'] for u in by_display.values() if u['id'] == uid), None)
        for d in list(days):
            if uname and (uname, d) in work_days:
                del days[d]
                conflicts += 1

    os.makedirs(args.out, exist_ok=True)
    HEADER = 'Дата\tВремя\tСотрудник\tПроект\tЧасы\tКомментарий\tID\n'

    def write_tsv(path, rows):
        with open(path, 'w', encoding='utf-8') as f:
            f.write(HEADER)
            for date_iso, username, project, hours in rows:
                y, m, d = date_iso.split('-')
                f.write(f"{d}.{m}.{y}\t\t{username}\t{project}\t{fmt_hours(hours)}\t\t{uuid.uuid4().hex[:12]}\n")

    matched_rows   = [(d, u, p, h) for d, u, p, h, m in reports if m]
    unmatched_rows = [(d, u, p, h) for d, u, p, h, m in reports if not m]
    write_tsv(os.path.join(args.out, 'reports_import.tsv'), matched_rows)
    if unmatched_rows:
        write_tsv(os.path.join(args.out, 'reports_unmatched.tsv'), unmatched_rows)

    json.dump(vacations, open(os.path.join(args.out, 'vacations_import.json'), 'w', encoding='utf-8'),
              ensure_ascii=False, indent=2)
    flags_plain = {uid: {d: dict(pr) for d, pr in days.items()} for uid, days in flags.items()}
    json.dump(flags_plain, open(os.path.join(args.out, 'report_flags_import.json'), 'w', encoding='utf-8'),
              ensure_ascii=False, indent=2)

    unmatched_names = sorted(canon[k] for k in canon if k not in matched_keys)
    if unmatched_names:
        with open(os.path.join(args.out, 'unmatched.txt'), 'w', encoding='utf-8') as f:
            f.write('\n'.join(unmatched_names))

    matched_users = sorted({u for _, u, _, _ in matched_rows})
    print(f"Источник:          лист «{SHEET}», месяцы {sorted(months_seen)}")
    print(f"Строк-отчётов:     {len(reports)}  (привязано: {len(matched_rows)} | без аккаунта: {len(unmatched_rows)})")
    print(f"Сотрудников сайта: {len(matched_users)} -> {matched_users}")
    print(f"Отпуск (дней, привязано к id): {sum(len(d) for d in vacations.values())}")
    print(f"Флагов переработка/отработка:  {sum(len(pr) for days in flags.values() for pr in days.values())}")
    print(f"Пропущено обедов (ячеек):      {skipped_lunch}")
    print(f"Конфликтов отпуск/часы (снято отпуска): {conflicts}")
    if unmatched_names:
        print(f"\nНе в users.json ({len(unmatched_names)}): {unmatched_names}")


if __name__ == '__main__':
    main()
